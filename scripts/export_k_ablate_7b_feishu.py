#!/usr/bin/env python3
"""Write the Feishu Excel copy of the 7B k-ablation table."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "results" / "reports" / "k_ablate_7b.json"
OUT = ROOT / "tables" / "firstwin_wait" / "k_ablate_7b_feishu.xlsx"

METHODS = ("full", "k2", "k3", "k4", "k5", "k6")
COMPARE = ("k2", "k3", "k4", "k5", "k6")
HEADERS = ["集", "n", "Full-CoT", "k=2", "k=3", "k=4", "k=5", "k=6"]

BOLD_IF = InlineFont(b=True, sz=11)
NORM_IF = InlineFont(b=False, sz=11)
HEAD_FONT = Font(bold=True, size=11)
BODY_FONT = Font(size=11)
OVER_FILL = PatternFill("solid", fgColor="FFF3E8")
HEAD_FILL = PatternFill("solid", fgColor="FFE8EEF4")
THIN = Border(
    left=Side(style="thin", color="FFD0D0D0"),
    right=Side(style="thin", color="FFD0D0D0"),
    top=Side(style="thin", color="FFD0D0D0"),
    bottom=Side(style="thin", color="FFD0D0D0"),
)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def acc_s(acc: float) -> str:
    return f"{acc:.2f}%"


def tok_s(tok: float) -> str:
    return str(int(round(tok)))


def rich_cell(acc: float, tok: float, acc_win: bool, tok_win: bool) -> CellRichText:
    return CellRichText(
        TextBlock(BOLD_IF if acc_win else NORM_IF, acc_s(acc)),
        TextBlock(NORM_IF, " / "),
        TextBlock(BOLD_IF if tok_win else NORM_IF, tok_s(tok)),
    )


def winners(parts: dict) -> tuple[set[str], set[str]]:
    accs = {key: parts[key]["acc"] for key in COMPARE}
    toks = {key: parts[key]["tok"] for key in COMPARE}
    max_acc = max(accs.values())
    min_tok = min(toks.values())
    return (
        {key for key, value in accs.items() if value == max_acc},
        {key for key, value in toks.items() if value == min_tok},
    )


def put_row(ws, widths, r, item, overall=False) -> None:
    acc_w, tok_w = winners(item)
    n = "—" if item.get("n") is None else item["n"]
    values = [item["dataset"], n]
    shown = [item["dataset"], str(n)]
    for key in METHODS:
        cell = item[key]
        acc_win = key in acc_w
        tok_win = key in tok_w
        values.append(rich_cell(cell["acc"], cell["tok"], acc_win, tok_win))
        shown.append(f"{acc_s(cell['acc'])} / {tok_s(cell['tok'])}")
    for col, value in enumerate(values, 1):
        cell = ws.cell(r, col, value)
        cell.font = BODY_FONT
        cell.border = THIN
        cell.alignment = LEFT if col == 1 else CENTER
        if overall:
            cell.fill = OVER_FILL
        widths[col - 1] = max(widths[col - 1], len(shown[col - 1]))


def main() -> None:
    payload = json.loads(SRC.read_text())
    rows = payload["cells"]
    wb = Workbook()
    ws = wb.active
    ws.title = "7B k ablation"
    ws.append(HEADERS)
    for col, _ in enumerate(HEADERS, 1):
        cell = ws.cell(1, col)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = THIN

    widths = [len(header) for header in HEADERS]
    r = 2
    for item in rows:
        put_row(ws, widths, r, item, overall=item.get("n") is None)
        r += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{r - 1}"
    ws.row_dimensions[1].height = 22
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            22, max(10, width * 1.15 + 2)
        )
    wb.save(OUT)
    print(f"wrote {OUT}")
    print(f"rows {r - 1}")


if __name__ == "__main__":
    main()
