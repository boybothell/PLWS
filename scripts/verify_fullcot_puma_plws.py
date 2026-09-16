#!/usr/bin/env python3
"""Independent recount of the main table Acc / token cells."""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from plws.matrix import deer_output_dir, job_rows
from plws.paths import PLWSPaths

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from report_fullcot_puma_plws import (  # noqa: E402
    OVERALL_EQ_LABEL,
    OVERALL_NW_LABEL,
    OVERALL_ONLY_DATASETS,
    MODELS,
    boxed_trial_tokens_upto,
    collect_shown,
    fmt_cell,
    overall_method_cells,
    collect_plws_scores,
    datasets_for,
    deer_seed_n,
    deer_seeds_for_cell,
    four_seed_cell_complete,
    load_json,
    load_trials_by_question,
    puma_delivery_path,
    puma_token,
    seeds_for_cell,
)

TABLE = ROOT / "tables" / "firstwin_wait" / "fullcot_puma_plws.md"
REPORT = ROOT / "results" / "reports" / "fullcot_puma_plws.json"
XLSX = ROOT / "tables" / "firstwin_wait" / "fullcot_puma_plws_feishu.xlsx"
CACHE = ROOT / "results" / "reports" / "cache" / "deer_aligned_v2_cells.json"


def mean(xs: list[float]) -> float:
    return sum(xs) / len(xs)


def recount(paths: PLWSPaths) -> dict:
    cells = {}
    issues: list[str] = []
    for model, zh in MODELS:
        for dataset, dszh, expect_n in datasets_for(model):
            allow_legacy = (dataset, dszh, expect_n) in OVERALL_ONLY_DATASETS
            if not four_seed_cell_complete(
                paths,
                model,
                dataset,
                expect_n,
                allow_legacy=allow_legacy,
            ):
                continue
            full_ok = full_tok = puma_ok = puma_tok = plws_ok = plws_tok = 0
            puma_del = puma_trial = plws_del = plws_trial = 0.0
            n = n_win = n_un = 0
            missing_trial = 0
            used_seeds = seeds_for_cell(
                paths, model, dataset, expect_n, allow_legacy=allow_legacy
            )
            want_n = deer_seed_n(expect_n) * len(used_seeds)
            for seed in used_seeds:
                official = {
                    int(r["question_idx"]): r
                    for r in load_json(paths.puma_statistics_path(model, dataset, seed))
                }
                delivery = load_json(puma_delivery_path(paths, model, dataset, seed))
                delivery_rows = (
                    delivery if isinstance(delivery, list) else delivery.get("rows", [])
                )
                if len(official) != len(delivery_rows):
                    issues.append(
                        f"{zh} {dszh} s{seed} official={len(official)} puma={len(delivery_rows)}"
                    )
                for row in official.values():
                    n += 1
                    full_ok += int(bool(row.get("original_correct")))
                    full_tok += float(row.get("original_tokens") or 0)
                for row in delivery_rows:
                    puma_ok += int(bool(row.get("compressed_correct")))
                    puma_tok += puma_token(row)
                    puma_del += float(row.get("compressed_tokens") or 0)
                    puma_trial += float(row.get("tokens_trial_answers") or 0)
                jobs = {
                    int(j["question_idx"]): j
                    for j in job_rows(paths, model, dataset, seed)
                    if j.get("uid")
                }
                scores = collect_plws_scores(
                    paths, model, dataset, seed, [], allow_legacy=allow_legacy
                )
                trials = {}
                tpath = paths.dense_trial_path(model, dataset, seed)
                if tpath.is_file():
                    trials = load_trials_by_question(tpath)
                for qid, info in official.items():
                    rec = scores.get((dataset, int(qid)))
                    if rec is not None:
                        n_win += 1
                        left = rec.get("left_step")
                        if left is None and qid in jobs:
                            left = jobs[qid].get("left_step")
                        if left is None:
                            issues.append(f"no left_step {model} {dataset} s{seed} q{qid}")
                            left = 0
                        trial = boxed_trial_tokens_upto(trials.get(int(qid), []), int(left))
                        if int(left) > 0 and not trials.get(int(qid)):
                            missing_trial += 1
                            issues.append(
                                f"no trials {model} {dataset} s{seed} q{qid} left={left}"
                            )
                        delivery = float(rec.get("n_think_tok") or 0) + float(
                            rec.get("n_ans_tok") or 0
                        )
                        plws_ok += int(bool(rec.get("new_gold_ok")))
                        plws_del += delivery
                        plws_trial += trial
                        plws_tok += delivery + trial
                    elif int(qid) not in jobs:
                        n_un += 1
                        plws_ok += int(bool(info.get("original_correct")))
                        plws_del += float(info.get("original_tokens") or 0)
                        plws_tok += float(info.get("original_tokens") or 0)
                    else:
                        issues.append(f"missing score {model} {dataset} s{seed} q{qid}")
            if n != want_n:
                issues.append(f"{zh} {dszh} n={n} expected {want_n}")
            deer = None
            deer_seeds = (
                ()
                if allow_legacy
                else deer_seeds_for_cell(
                    paths, model, dataset, expect_n, used_seeds
                )
            )
            if deer_seeds:
                d_ok = d_tok = d_del = d_trial = 0.0
                d_n = 0
                for seed in deer_seeds:
                    path = deer_output_dir(paths, model, dataset, seed) / "deer.jsonl"
                    for line in path.read_text().splitlines():
                        if not line.strip():
                            continue
                        row = json.loads(line)
                        delivery = float(row.get("delivery_tokens") or 0)
                        if delivery <= 0:
                            delivery = float(row.get("thinking_tokens") or 0) + float(
                                row.get("answer_tokens") or 0
                            )
                        trial = float(row.get("num_trial_answer_tokens") or 0)
                        d_n += 1
                        d_del += delivery
                        d_trial += trial
                        d_tok += delivery + trial
                deer = {
                    "n": d_n,
                    "tok": d_tok / d_n,
                    "tok_delivery": d_del / d_n,
                    "tok_trial": d_trial / d_n,
                }
            cells[(zh, dszh)] = {
                "n": n,
                "seeds": list(used_seeds),
                "deer_seeds": list(deer_seeds),
                "windowed": n_win,
                "unwindowed": n_un,
                "missing_trial": missing_trial,
                "full": {"acc": 100.0 * full_ok / n, "tok": full_tok / n},
                "puma": {
                    "acc": 100.0 * puma_ok / n,
                    "tok": puma_tok / n,
                    "tok_delivery": puma_del / n,
                    "tok_trial": puma_trial / n,
                },
                "plws": {
                    "acc": 100.0 * plws_ok / n,
                    "tok": plws_tok / n,
                    "tok_delivery": plws_del / n,
                    "tok_trial": plws_trial / n,
                },
                "deer": deer,
            }
    return {"cells": cells, "issues": issues}


def close(a: float, b: float, acc: bool = False) -> bool:
    if acc:
        return abs(round(a, 2) - round(b, 2)) < 1e-9 or abs(a - b) < 0.005
    return abs(round(a) - round(b)) <= 0 or abs(a - b) < 0.51


def main() -> None:
    paths = PLWSPaths.discover(ROOT)
    report = json.loads(REPORT.read_text())
    table = TABLE.read_text()
    cache = json.loads(CACHE.read_text()) if CACHE.is_file() else {}
    print("recounting from artifacts...", flush=True)
    got = recount(paths)
    mismatches: list[str] = []
    for item in got["issues"]:
        mismatches.append(f"ISSUE {item}")
    by_report = {(row["model"], row["dataset"]): row for row in report["cells"]}
    for model, zh in MODELS:
        accs = defaultdict(list)
        toks = defaultdict(list)
        ns = defaultdict(list)
        for dataset, dszh, expect_n in datasets_for(model):
            hidden = (dataset, dszh, expect_n) in OVERALL_ONLY_DATASETS
            if (zh, dszh) not in got["cells"]:
                if (zh, dszh) in by_report:
                    mismatches.append(f"{zh} {dszh} incomplete but still in table")
                continue
            cell = got["cells"][(zh, dszh)]
            ref = by_report[(zh, dszh)]
            if cell["n"] != ref["n"]:
                mismatches.append(f"{zh} {dszh} n {cell['n']} vs table {ref['n']}")
            if cell["seeds"] != ref.get("seeds", cell["seeds"]):
                mismatches.append(
                    f"{zh} {dszh} seeds {cell['seeds']} vs table {ref.get('seeds')}"
                )
            for key in ("full", "puma", "plws"):
                g, r = cell[key], ref[key]
                accs[key].append(g["acc"])
                toks[key].append(g["tok"])
                ns[key].append(ref["n"])
                if abs(g["acc"] - r["acc"]) > 0.005:
                    mismatches.append(
                        f"{zh} {dszh} {key} acc {g['acc']:.4f} vs table {r['acc']}"
                    )
                if abs(g["tok"] - r["tok"]) > 0.51:
                    mismatches.append(
                        f"{zh} {dszh} {key} tok {g['tok']:.3f} vs table {r['tok']}"
                    )
                shown = f"{g['acc']:.2f}% / {round(g['tok'])}"
                plain = table.replace("**", "")
                if shown not in plain and f"{g['acc']:.2f}% / {int(g['tok'])}" not in plain:
                    # table uses round()
                    pass
                if shown not in plain:
                    mismatches.append(f"{zh} {dszh} {key} md missing {shown}")
                if key != "full" and g["tok_trial"] < 0:
                    mismatches.append(f"{zh} {dszh} {key} negative trial")
                if key == "puma" and g["tok_trial"] <= 0:
                    mismatches.append(f"{zh} {dszh} PUMA trial is 0")
            deer_g, deer_r = cell["deer"], ref["deer"]
            if hidden and (deer_r is not None or deer_g is not None):
                mismatches.append(f"{zh} {dszh} overall-only DEER should be empty")
            if deer_r is None:
                if deer_g is not None:
                    mismatches.append(f"{zh} {dszh} DEER should be empty")
                if "|  |" not in table.split(f"| {zh} | {dszh} |", 1)[-1].split("\n", 1)[0] and deer_g is None:
                    # empty deer cell in markdown is " |  | " or " |  |"
                    pass
            else:
                if deer_g is None:
                    mismatches.append(f"{zh} {dszh} DEER missing in recount")
                    continue
                if deer_g["n"] != deer_seed_n(expect_n) * len(cell["deer_seeds"]):
                    mismatches.append(f"{zh} {dszh} DEER n {deer_g['n']}")
                if abs(deer_g["tok"] - deer_r["tok"]) > 0.51:
                    mismatches.append(
                        f"{zh} {dszh} DEER tok {deer_g['tok']:.3f} vs table {deer_r['tok']}"
                    )
                shown = f"{deer_r['acc']:.2f}% / {deer_r['tok']}"
                if shown not in table.replace("**", ""):
                    mismatches.append(f"{zh} {dszh} DEER md missing {shown}")
                tag = {
                    "7B": "r1_7b",
                    "Nemotron": "nemotron_8b",
                    "14B": "r1_14b",
                    "1.5B": "r1_1p5b",
                    "Llama-8B": "r1_llama_8b",
                    "R1-32B": "r1_32b",
                    "4B": "qwen3_4b",
                    "8B": "qwen3_8b",
                    "30B": "qwen3_30b_a3b",
                }[zh]
                hit = cache.get(f"{tag}/{dataset}")
                if not hit or abs(hit["acc"] - deer_r["acc"]) > 0.005:
                    mismatches.append(
                        f"{zh} {dszh} DEER acc cache {None if not hit else hit['acc']} vs table {deer_r['acc']}"
                    )
                accs["deer"].append(deer_r["acc"])
                toks["deer"].append(deer_g["tok"])
                ns["deer"].append(deer_g["n"])
            print(
                f"{zh:8} {dszh:14} full={cell['full']['acc']:.2f}/{round(cell['full']['tok'])} "
                f"puma={cell['puma']['acc']:.2f}/{round(cell['puma']['tok'])} "
                f"trialP={cell['puma']['tok_trial']:.1f} "
                f"plws={cell['plws']['acc']:.2f}/{round(cell['plws']['tok'])} "
                f"trialL={cell['plws']['tok_trial']:.1f} win={cell['windowed']} un={cell['unwindowed']} "
                f"deerTok={None if deer_g is None else round(deer_g['tok'])} "
                f"missTrialQ={cell['missing_trial']}"
            )
        model_rows = [
            row for row in report["cells"] if row["model"] == zh
        ]
        if not model_rows:
            continue
        equal, weighted = overall_method_cells(collect_shown(model_rows))
        plain = table.replace("**", "")
        for label, cells in (
            (OVERALL_EQ_LABEL, equal),
            (OVERALL_NW_LABEL, weighted),
        ):
            for key in ("full", "puma", "plws", "deer"):
                cell = cells.get(key)
                if cell is None:
                    continue
                shown = fmt_cell(cell["acc"], cell["tok"])
                if shown not in plain:
                    mismatches.append(f"{zh} {label} {key} md missing {shown}")

    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError:
        print("xlsx skipped (no openpyxl)")
        if mismatches:
            print("MISMATCHES")
            for item in mismatches:
                print(" ", item)
            raise SystemExit(1)
        print("OK Acc/token cells match artifacts, markdown, and JSON")
        return
    wb = load_workbook(XLSX, rich_text=True)
    ws = wb.active
    xrows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7, values_only=False))
    expect_xlsx = sum(
        sum(
            1
            for dataset, _dszh, expect_n in datasets_for(model)
            if four_seed_cell_complete(
                paths,
                model,
                dataset,
                expect_n,
                allow_legacy=(dataset, _dszh, expect_n) in OVERALL_ONLY_DATASETS,
            )
        )
        + 2
        for model, _zh in MODELS
    )
    if len(xrows) != expect_xlsx:
        mismatches.append(f"xlsx rows {len(xrows)} expected {expect_xlsx}")
    names = {
        "7B": "DeepSeek-R1-Distill-Qwen-7B",
        "Nemotron": "Llama-3.1-Nemotron-Nano-8B-v1",
        "14B": "DeepSeek-R1-Distill-Qwen-14B",
        "1.5B": "DeepSeek-R1-Distill-Qwen-1.5B",
        "Llama-8B": "DeepSeek-R1-Distill-Llama-8B",
        "R1-32B": "DeepSeek-R1-Distill-Qwen-32B",
        "4B": "Qwen3-4B",
        "8B": "Qwen3-8B",
        "30B": "Qwen3-30B-A3B-Thinking-2507",
    }
    xi = 0
    for model, zh in MODELS:
        row_specs = [
            item
            for item in datasets_for(model)
            if four_seed_cell_complete(
                paths,
                model,
                item[0],
                item[2],
                allow_legacy=item in OVERALL_ONLY_DATASETS,
            )
        ] + [
            ("overall", OVERALL_EQ_LABEL, None),
            ("overall_nw", OVERALL_NW_LABEL, None),
        ]
        for dataset, dszh, expect_n in row_specs:
            cell = xrows[xi]
            xi += 1
            name = str(cell[0].value)
            ds = str(cell[1].value)
            if name != names[zh] or (ds != dszh and not ds.startswith("Overall")):
                mismatches.append(f"xlsx label {name} {ds} vs {zh} {dszh}")
            if dszh.startswith("Overall"):
                continue
            ref = by_report[(zh, dszh)]
            def txt(c):
                v = c.value
                return "" if v is None else str(v)
            want = [
                f"{ref['full']['acc']:.2f}% / {ref['full']['tok']}",
                f"{ref['puma']['acc']:.2f}% / {ref['puma']['tok']}",
                "" if ref["deer"] is None else f"{ref['deer']['acc']:.2f}% / {ref['deer']['tok']}",
                f"{ref['plws']['acc']:.2f}% / {ref['plws']['tok']}",
            ]
            gotv = [txt(cell[3]), txt(cell[4]), txt(cell[5]), txt(cell[6])]
            if gotv != want:
                mismatches.append(f"xlsx {zh} {dszh} {gotv} vs {want}")

    if got["issues"] or mismatches:
        print("MISMATCHES")
        for item in mismatches:
            print(" ", item)
        raise SystemExit(1)
    print("OK all Acc/token cells match artifacts, markdown, JSON, and xlsx")


if __name__ == "__main__":
    main()
