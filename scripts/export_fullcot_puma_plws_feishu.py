#!/usr/bin/env python3
"""Write the Feishu Excel copy of the Full-CoT / PUMA / DEER / 窗后压 table.

Each method cell bolds the highest Acc and the lowest token among Full-CoT,
PUMA, DEER, and 窗后压. Empty DEER cells are skipped.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # unittest / okay-budget venv has no openpyxl
    Workbook = CellRichText = TextBlock = InlineFont = None
    Alignment = Border = Font = PatternFill = Side = get_column_letter = None

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "results" / "reports" / "fullcot_puma_plws.json"
OUT = ROOT / "tables" / "firstwin_wait" / "fullcot_puma_plws_feishu.xlsx"
sys.path.insert(0, str(ROOT / "scripts"))
from report_fullcot_puma_plws import (  # noqa: E402
    COMPARE,
    MODELS,
    OVERALL_EQ_LABEL,
    OVERALL_NW_LABEL,
    collect_shown,
    datasets_for,
    overall_method_cells,
    shown_row,
    winners,
)

NAMES = {
    "7B": "DeepSeek-R1-Distill-Qwen-7B",
    "Nemotron": "Llama-3.1-Nemotron-Nano-8B-v1",
    "14B": "DeepSeek-R1-Distill-Qwen-14B",
    "1.5B": "DeepSeek-R1-Distill-Qwen-1.5B",
    "Llama-8B": "DeepSeek-R1-Distill-Llama-8B",
    "R1-32B": "DeepSeek-R1-Distill-Qwen-32B",
    "4B": "Qwen3-4B",
    "8B": "Qwen3-8B",
    "30B": "Qwen3-30B-A3B-Thinking-2507",
}
ORDER = [zh for _tag, zh in MODELS]
METHODS = COMPARE
HEADERS = ["模型", "集", "n", "Full-CoT (原始 CoT)", "PUMA", "DEER", "窗后压"]
THREE_HEADERS = ["模型", "集", "n", "入选 seed", "Full-CoT (原始 CoT)", "PUMA", "DEER", "窗后压"]

if InlineFont is not None:
    BOLD_IF = InlineFont(b=True, sz=11)
    NORM_IF = InlineFont(b=False, sz=11)
    HEAD_FONT = Font(bold=True, size=11)
    BODY_FONT = Font(size=11)
    OVER_FILL = PatternFill("solid", fgColor="FFF3E8")
    HEAD_FILL = PatternFill("solid", fgColor="FFE8EEF4")
    THIN = Border(
        left=Side(style="thin", color="FFD0D0D0"),
        right=Side(style="thin", color="FFD0D0D0"),
        top=Side(style="thin", color="FFD0D0D0"),
        bottom=Side(style="thin", color="FFD0D0D0"),
    )
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT = Alignment(horizontal="left", vertical="center")
else:
    BOLD_IF = NORM_IF = HEAD_FONT = BODY_FONT = None
    OVER_FILL = HEAD_FILL = THIN = CENTER = LEFT = None


def acc_s(acc: float) -> str:
    return f"{acc:.2f}%"


def tok_s(tok: float) -> str:
    return str(int(round(tok)))


def rich_cell(acc: float, tok: float, acc_win: bool, tok_win: bool) -> CellRichText:
    return CellRichText(
        TextBlock(BOLD_IF if acc_win else NORM_IF, acc_s(acc)),
        TextBlock(NORM_IF, " / "),
        TextBlock(BOLD_IF if tok_win else NORM_IF, tok_s(tok)),
    )


def method_value(parts: dict, key: str, acc_w: set[str], tok_w: set[str]):
    cell = parts.get(key)
    if cell is None:
        return "", ""
    return (
        rich_cell(cell["acc"], cell["tok"], key in acc_w, key in tok_w),
        f"{acc_s(cell['acc'])} / {tok_s(cell['tok'])}",
    )


def put_row(ws, widths, r, model, dataset, n, parts, overall=False, extra=None) -> None:
    acc_w, tok_w = winners(parts)
    extras = list(extra or [])
    values = [model, dataset, n, *extras]
    shown = [model, dataset, str(n), *[str(item) for item in extras]]
    for key in METHODS:
        value, text = method_value(parts, key, acc_w, tok_w)
        values.append(value)
        shown.append(text or "00.00% / 00000")
    for col, value in enumerate(values, 1):
        cell = ws.cell(r, col, value)
        cell.font = BODY_FONT
        cell.border = THIN
        cell.alignment = LEFT if col <= 2 else CENTER
        if overall:
            cell.fill = OVER_FILL
        widths[col - 1] = max(widths[col - 1], len(shown[col - 1]))


def _style_header(ws, headers) -> None:
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = THIN


def _autosize(ws, widths, last_row: int) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{last_row}"
    ws.row_dimensions[1].height = 22
    for i, width in enumerate(widths, 1):
        factor = 1.2 if i == 1 else 1.15
        ws.column_dimensions[get_column_letter(i)].width = min(
            48, max(10, width * factor + 2)
        )


def _seeds_text(row: dict | None) -> str:
    if not row or not row.get("seeds"):
        return "—"
    return "/".join(str(seed) for seed in row["seeds"])


def fill_method_sheet(ws, payload: dict, title: str, *, with_seeds: bool = False) -> int:
    ws.title = title
    headers = THREE_HEADERS if with_seeds else HEADERS
    ws.append(headers)
    _style_header(ws, headers)
    by_model: dict[str, list] = defaultdict(list)
    for row in payload["cells"]:
        by_model[row["model"]].append(row)
    widths = [len(header) for header in headers]
    r = 2
    for tag, model in MODELS:
        rows = {item["dataset"]: item for item in by_model.get(model, [])}
        if not rows:
            continue
        shown = {key: [] for key in METHODS}
        name = NAMES[model]
        kept = []
        for _dataset, dszh, _n in datasets_for(tag):
            item = rows.get(dszh)
            if item is None:
                continue
            kept.append(item)
            put_row(
                ws,
                widths,
                r,
                name,
                item["dataset"],
                item["n"],
                item,
                extra=[_seeds_text(item)] if with_seeds else None,
            )
            for key in METHODS:
                item_row = shown_row(item, key)
                if item_row is not None:
                    shown[key].append(item_row)
            r += 1
        equal, weighted = overall_method_cells(shown)
        put_row(
            ws,
            widths,
            r,
            name,
            OVERALL_EQ_LABEL,
            "—",
            equal,
            overall=True,
            extra=["—"] if with_seeds else None,
        )
        r += 1
        put_row(
            ws,
            widths,
            r,
            name,
            OVERALL_NW_LABEL,
            weighted["full"]["n"] if kept else "—",
            weighted,
            overall=True,
            extra=["—"] if with_seeds else None,
        )
        r += 1
    _autosize(ws, widths, r - 1)
    return r - 1


def write_xlsx(src: Path, out: Path) -> None:
    if Workbook is None:
        raise SystemExit("openpyxl is required to write the Feishu xlsx")
    payload = json.loads(src.read_text())
    wb = Workbook()
    rows = fill_method_sheet(wb.active, payload, "Full-CoT PUMA DEER 窗后压")
    wb.save(out)
    print(f"wrote {out}")
    print(f"rows {rows}")


COMPARE_HEADERS = [
    "模型",
    "集",
    "入选 seed",
    "n 3seed",
    "n 5seed",
    "Full-CoT 3seed",
    "Full-CoT 5seed",
    "PUMA 3seed",
    "PUMA 5seed",
    "窗后压 3seed",
    "窗后压 5seed",
]


def put_compare_row(ws, widths, r, model, dataset, three, five, overall=False) -> None:
    values = [model, dataset, "—" if overall else _seeds_text(three)]
    shown = [model, dataset, str(values[-1])]
    values.append("—" if overall or three is None else three["n"])
    shown.append(str(values[-1]))
    values.append("—" if overall or five is None else five["n"])
    shown.append(str(values[-1]))
    for key in ("full", "puma", "plws"):
        left = None if three is None else three.get(key)
        right = None if five is None else five.get(key)
        acc_w3 = acc_w5 = tok_w3 = tok_w5 = False
        if left and right:
            acc_w3 = left["acc"] >= right["acc"]
            acc_w5 = right["acc"] >= left["acc"]
            tok_w3 = left["tok"] <= right["tok"]
            tok_w5 = right["tok"] <= left["tok"]
        elif left:
            acc_w3 = tok_w3 = True
        elif right:
            acc_w5 = tok_w5 = True
        for cell, acc_w, tok_w in ((left, acc_w3, tok_w3), (right, acc_w5, tok_w5)):
            if cell is None:
                values.append("")
                shown.append("00.00% / 00000")
            else:
                values.append(rich_cell(cell["acc"], cell["tok"], acc_w, tok_w))
                shown.append(f"{acc_s(cell['acc'])} / {tok_s(cell['tok'])}")
    for col, value in enumerate(values, 1):
        cell = ws.cell(r, col, value)
        cell.font = BODY_FONT
        cell.border = THIN
        cell.alignment = LEFT if col <= 2 else CENTER
        if overall:
            cell.fill = OVER_FILL
        widths[col - 1] = max(widths[col - 1], len(shown[col - 1]))


def _overall_from_cells(rows: list[dict], *, weighted: bool = False) -> dict | None:
    if not rows:
        return None
    equal, n_weighted = overall_method_cells(collect_shown(rows))
    out = n_weighted if weighted else equal
    out["n"] = "—" if not weighted else out["full"]["n"]
    out["dataset"] = OVERALL_NW_LABEL if weighted else OVERALL_EQ_LABEL
    return out


def write_3seed_xlsx(three_src: Path, five_src: Path, out: Path) -> None:
    if Workbook is None:
        raise SystemExit("openpyxl is required to write the Feishu xlsx")
    three = json.loads(three_src.read_text())
    five = json.loads(five_src.read_text())
    wb = Workbook()
    fill_method_sheet(wb.active, three, "三 seed", with_seeds=True)
    cmp = wb.create_sheet("对照五 seed")
    cmp.append(COMPARE_HEADERS)
    _style_header(cmp, COMPARE_HEADERS)
    widths = [len(header) for header in COMPARE_HEADERS]
    by3 = defaultdict(list)
    by5 = defaultdict(list)
    for row in three["cells"]:
        by3[row["model"]].append(row)
    for row in five["cells"]:
        by5[row["model"]].append(row)
    r = 2
    for tag, model in MODELS:
        rows3 = {item["dataset"]: item for item in by3.get(model, [])}
        rows5 = {item["dataset"]: item for item in by5.get(model, [])}
        names = []
        for _dataset, dszh, _n in datasets_for(tag):
            if dszh in rows3 or dszh in rows5:
                names.append(dszh)
        if not names:
            continue
        name = NAMES[model]
        kept3 = []
        kept5 = []
        for dszh in names:
            item3 = rows3.get(dszh)
            item5 = rows5.get(dszh)
            put_compare_row(cmp, widths, r, name, dszh, item3, item5)
            if item3 is not None:
                kept3.append(item3)
            if item5 is not None:
                kept5.append(item5)
            r += 1
        put_compare_row(
            cmp,
            widths,
            r,
            name,
            OVERALL_EQ_LABEL,
            _overall_from_cells(kept3),
            _overall_from_cells(kept5),
            overall=True,
        )
        r += 1
        put_compare_row(
            cmp,
            widths,
            r,
            name,
            OVERALL_NW_LABEL,
            _overall_from_cells(kept3, weighted=True),
            _overall_from_cells(kept5, weighted=True),
            overall=True,
        )
        r += 1
    _autosize(cmp, widths, r - 1)
    cmp.sheet_properties.tabColor = "FFF3E8"
    note = wb.create_sheet("说明")
    note["A1"] = "每一格从 42/0/1/123 自选三个 seed，不是全局共用一组。"
    note["A2"] = three.get("rule") or ""
    note["A3"] = "对照页同一格加粗：Acc 更高、token 更低；并列都加粗。"
    note["A4"] = "Llama-8B AIME26 / AMC23 四 seed 未齐，两页都不进。"
    note.column_dimensions["A"].width = 80
    wb.save(out)
    print(f"wrote {out}")
    print("3seed per-cell pick")


MEAN3_SRC = ROOT / "results" / "reports" / "fullcot_puma_plws_mean3.json"
MEAN3_OUT = ROOT / "tables" / "firstwin_wait" / "fullcot_puma_plws_mean3_feishu.xlsx"
MEAN3_HEADERS = ["模型", "集", "n", "Full-CoT", "PUMA", "DEER", "窗后压"]


def _tr_s(full: dict | None, cell: dict) -> str:
    if full is None:
        return "—"
    full_tok = float(full["tok"])
    if full_tok <= 0:
        return "0.0%"
    return f"{100.0 * (full_tok - float(cell['tok'])) / full_tok:.1f}%"


def rich_mean3_cell(cell: dict, full: dict | None, *, kind: str) -> CellRichText:
    acc = f"{cell['acc']:.2f}%"
    tok = str(int(round(cell["tok"])))
    blocks = [
        TextBlock(BOLD_IF if cell.get("acc_win") else NORM_IF, acc),
        TextBlock(NORM_IF, " / "),
        TextBlock(BOLD_IF if cell.get("tok_win") else NORM_IF, tok),
        TextBlock(NORM_IF, " / "),
    ]
    if kind == "full" or full is None:
        blocks.append(TextBlock(NORM_IF, "—"))
    else:
        blocks.append(
            TextBlock(BOLD_IF if cell.get("tr_win") else NORM_IF, _tr_s(full, cell))
        )
    return CellRichText(*blocks)


def write_mean3_xlsx(src: Path = MEAN3_SRC, out: Path = MEAN3_OUT) -> None:
    if Workbook is None:
        raise SystemExit("openpyxl is required to write the Feishu xlsx")
    sys.path.insert(0, str(ROOT / "scripts"))
    from report_fullcot_puma_plws_mean3 import (
        COMPARE,
        DATASETS,
        MODELS,
        OVERALL_EQ,
        overall_parts,
    )

    payload = json.loads(src.read_text())
    by_model: dict[str, list[dict]] = defaultdict(list)
    for row in payload["cells"]:
        by_model[row["model"]].append(row)
    wb = Workbook()
    ws = wb.active
    ws.title = "mean@3"
    ws.append(MEAN3_HEADERS)
    _style_header(ws, MEAN3_HEADERS)
    widths = [len(header) for header in MEAN3_HEADERS]
    r = 2
    for _tag, zh in MODELS:
        rows = [row for row in by_model.get(zh, [])]
        by_ds = {row["dataset_id"]: row for row in rows}
        ordered = [by_ds[ds] for ds, _dszh, _n in DATASETS if ds in by_ds]
        if not ordered:
            continue
        name = NAMES[zh]
        for row in ordered:
            values = [name, row["dataset"], row["n"]]
            shown = [name, row["dataset"], str(row["n"])]
            for key in COMPARE:
                cell = row.get(key)
                if cell is None:
                    values.append("")
                    shown.append("未齐")
                else:
                    values.append(rich_mean3_cell(cell, row.get("full"), kind=key))
                    shown.append(
                        f"{cell['acc']:.2f}% / {int(round(cell['tok']))} / "
                        f"{'—' if key == 'full' else _tr_s(row.get('full'), cell)}"
                    )
            for col, value in enumerate(values, 1):
                item = ws.cell(r, col, value)
                item.font = BODY_FONT
                item.border = THIN
                item.alignment = LEFT if col <= 2 else CENTER
                widths[col - 1] = max(widths[col - 1], len(shown[col - 1]))
            r += 1
        if len(ordered) == len(DATASETS) and all(
            all(row["counts"][key] == 3 for key in ("full", "puma", "plws"))
            for row in ordered
        ):
            parts = overall_parts(ordered, weighted=False)
            values = [name, OVERALL_EQ, "—"]
            shown = [name, OVERALL_EQ, "—"]
            for key in COMPARE:
                cell = parts.get(key)
                if cell is None:
                    values.append("")
                    shown.append("")
                else:
                    values.append(rich_mean3_cell(cell, parts.get("full"), kind=key))
                    shown.append(
                        f"{cell['acc']:.2f}% / {int(round(cell['tok']))} / "
                        f"{'—' if key == 'full' else _tr_s(parts.get('full'), cell)}"
                    )
            for col, value in enumerate(values, 1):
                item = ws.cell(r, col, value)
                item.font = BODY_FONT
                item.border = THIN
                item.alignment = LEFT if col <= 2 else CENTER
                item.fill = OVER_FILL
                widths[col - 1] = max(widths[col - 1], len(shown[col - 1]))
            r += 1
    _autosize(ws, widths, r - 1)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")
    print(f"rows {r - 1}")


def main() -> None:
    write_xlsx(SRC, OUT)


if __name__ == "__main__":
    main()
